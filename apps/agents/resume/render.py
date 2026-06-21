"""Render a tailored resume (plain structured text from tailor.py) into an
ATS-safe .docx -- single column, standard headers, no tables/graphics.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.shared import Pt

SECTION_HEADERS = {"summary", "skills", "experience", "projects", "education", "coverage gaps"}

BULLET_SOURCE_TAG_RE = re.compile(r"\s*\[fact_[^\]]+\]\s*$")


def _strip_source_tags(line: str) -> str:
    """Remove the [fact_id] trace tags before producing the human-facing document."""
    return BULLET_SOURCE_TAG_RE.sub("", line).strip()


def render_docx(tailored_text: str, output_path: str | Path) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    for raw_line in tailored_text.splitlines():
        line = raw_line.strip()
        if not line:
            continue

        lowered = line.lower().rstrip(":")
        if lowered in SECTION_HEADERS:
            doc.add_heading(line.rstrip(":").title(), level=2)
            continue

        clean_line = _strip_source_tags(line.lstrip("•-*").strip())
        if not clean_line:
            continue

        doc.add_paragraph(clean_line, style="List Bullet" if line.startswith(("•", "-", "*")) else "Normal")

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
